"""Bounded, deterministic .xlsx-to-ledger parser. No I/O outside supplied bytes."""

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

from app.v2.imports.models import Direction, ExcludedRow, ParseError, ParsedRow, ParsedWorkbook, RawCell

MAX_COMPRESSED_BYTES = 10_485_760
MAX_EXPANDED_BYTES = 104_857_600
MAX_SHEETS = 50
MAX_POPULATED_ROWS = 100_000
MAX_SCANNED_ROWS = 200_000
MONEY_UNIT = Decimal("0.01")

HEADER_WORDS = (
    "date", "description", "descn", "particulars", "details", "narration", "amount", "amt",
    "value", "rate", "qty", "quantity", "weight", "bags", "party", "supplier", "customer",
    "advance", "paid", "balance", "debit", "credit", "revenue", "expense", "remarks",
)
DESCRIPTION_HEADERS = (
    "description", "descn", "particulars", "details", "narration", "item", "product",
    "purpose", "party", "supplier", "customer", "name",
)
AMOUNT_HEADERS = ("amount", "amt", "total amount", "value", "revenue", "expense", "total paid", "paid", "advance")
TOTAL_LABELS = {"total", "grand total", "closing balance", "opening balance"}
MONEY_PATTERN = re.compile(r"^-?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?$")


def _normalize(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def _empty(value: object) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def _raw_cells(values: tuple[object, ...]) -> dict[str, RawCell]:
    raw: dict[str, RawCell] = {}
    for index, value in enumerate(values, start=1):
        if _empty(value):
            continue
        if isinstance(value, (date, datetime)):
            safe_value: RawCell = value.isoformat()
        elif isinstance(value, float):
            safe_value = str(value)
        elif isinstance(value, (str, int, bool)):
            safe_value = value
        else:
            safe_value = str(value)
        raw[get_column_letter(index)] = safe_value
    return raw


def _check_archive(content: bytes) -> None:
    if not content or len(content) > MAX_COMPRESSED_BYTES:
        raise ParseError("file_size", "The workbook must be between 1 byte and 10 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ParseError("unsupported_file", "Upload a valid .xlsx workbook.")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise ParseError("unsupported_file", "Macro-enabled workbooks are not supported.")
            expanded_size = 0
            for member in archive.infolist():
                if member.flag_bits & 1:
                    raise ParseError("encrypted_file", "Encrypted workbooks are not supported.")
                expanded_size += member.file_size
                if expanded_size > MAX_EXPANDED_BYTES:
                    raise ParseError("expanded_size", "The workbook expands beyond the 100 MB safety limit.")
    except BadZipFile as error:
        raise ParseError("unsupported_file", "Upload a valid .xlsx workbook.") from error


def _header_index(rows: list[tuple[int, tuple[object, ...]]]) -> int | None:
    best_index = None
    best_score = 0
    for index, (_number, row) in enumerate(rows[:30]):
        values = [value for value in row if not _empty(value)]
        if len(values) < 2:
            continue
        matches = sum(any(word == _normalize(value) or word in _normalize(value) for word in HEADER_WORDS)
                      for value in values)
        score = matches * 10 + min(len(values), 8)
        if matches >= 2 and score > best_score:
            best_index = index
            best_score = score
    return best_index


def _column(headers: tuple[object, ...], candidates: tuple[str, ...], sheet: str) -> int | None:
    normalized = [_normalize(value) for value in headers]
    exact = [index for index, value in enumerate(normalized) if value and value in candidates]
    matches = exact or [index for index, value in enumerate(normalized)
                        if value and any(candidate in value for candidate in candidates)]
    if len(matches) > 1:
        raise ParseError("ambiguous_columns", f"{sheet}: multiple columns match {', '.join(candidates[:2])}.")
    return matches[0] if matches else None


def _date(value: object) -> date | None:
    parsed: date | None = None
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    elif isinstance(value, str):
        text = value.strip()
        match = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", text)
        if match:
            year = int(match[3])
            if year < 100:
                year += 2000
            try:
                parsed = date(year, int(match[2]), int(match[1]))
            except ValueError:
                return None
        else:
            match = re.fullmatch(r"(\d{1,2})\.(\d{2})(\d{4})", text)
            if match:
                try:
                    parsed = date(int(match[3]), int(match[2]), int(match[1]))
                except ValueError:
                    return None
            else:
                match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
                if match:
                    try:
                        parsed = date(int(match[1]), int(match[2]), int(match[3]))
                    except ValueError:
                        return None
    return parsed if parsed and 2000 <= parsed.year <= 2100 else None


def _amount(value: object, location: str) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        text = value.strip().replace("₹", "").replace("$", "").strip()
        if text.startswith("(") and text.endswith(")"):
            text = f"-{text[1:-1]}"
        if not MONEY_PATTERN.fullmatch(text):
            return None
        text = text.replace(",", "")
    elif isinstance(value, (int, float, Decimal)):
        text = str(value)
    else:
        return None
    try:
        amount = Decimal(text)
        rounded = amount.quantize(MONEY_UNIT)
    except (InvalidOperation, ValueError) as error:
        raise ParseError("invalid_amount", f"{location}: amount is not a finite currency value.") from error
    if not amount.is_finite() or amount != rounded:
        raise ParseError("amount_precision", f"{location}: amount needs at most two decimal places.")
    return rounded


def _sheet_kind(name: str, descriptions: list[str]) -> Direction | None:
    sheet = _normalize(name)
    if re.search(r"pending|debt|due|balance|stock|summary", sheet):
        return None
    if re.search(r"breakup|mixed|reconciliation", sheet):
        return None
    if re.search(r"revenue|income|sales|receipts?", sheet):
        return "income"
    if re.search(r"expense|cost|purchase|advance|payment|paddy|transport", sheet):
        return "expense"
    sample = _normalize(" ".join(descriptions[:120]))
    expense_words = ("repair", "charges", "petrol", "diesel", "labour", "wages", "advance",
                     "purchase", "paid", "cost", "emi", "food", "transport")
    income_words = ("revenue", "sale", "receipt", "received", "weighbridge")
    expense_score = sum(word in sample for word in expense_words)
    income_score = sum(word in sample for word in income_words)
    if expense_score >= 2 and expense_score > income_score:
        return "expense"
    if income_score >= 2 and income_score > expense_score:
        return "income"
    return None


def _category(description: str, direction: Direction) -> str:
    text = _normalize(description)
    if direction == "income":
        for pattern, category in ((r"bran", "Rice bran"), (r"husk|vuka", "Rice husk"),
                                  (r"broken|nuka", "Broken rice"), (r"weigh", "Weighbridge"),
                                  (r"paddy", "Paddy sales")):
            if re.search(pattern, text):
                return category
        return "Other income"
    for pattern, category in ((r"paddy|farmer|supplier", "Paddy procurement"),
                              (r"transport|vehicle|driver|diesel|petrol|freight", "Transport & fuel"),
                              (r"labour|labor|wage|salary|worker", "Labour & wages"),
                              (r"repair|bearing|motor|maintenance|machine", "Repairs & maintenance"),
                              (r"officer|licen[cs]e|registration|gst|commission|bank|\bbg\b", "Compliance & bank"),
                              (r"electric|power|tea|food|bag|wire|gunny|office", "Mill operations"),
                              (r"advance|payment|paid", "Advances & payments")):
        if re.search(pattern, text):
            return category
    return "Other expenses"


def _row_value(row: tuple[object, ...], index: int | None) -> object:
    return row[index] if index is not None and index < len(row) else None


def _parse_sheet(name: str, rows: list[tuple[int, tuple[object, ...], frozenset[int]]]
                 ) -> tuple[list[ParsedRow], list[ExcludedRow]]:
    included: list[ParsedRow] = []
    excluded: list[ExcludedRow] = []
    values_only = [(number, values) for number, values, _formulas in rows]
    header_index = _header_index(values_only)
    reference_sheet = re.search(r"pending|debt|due|balance|stock|summary", _normalize(name)) is not None
    if reference_sheet:
        return [], [ExcludedRow(name, number, _raw_cells(values), "reference_or_summary_sheet")
                    for number, values, _formulas in rows]
    if header_index is None:
        raise ParseError("missing_header", f"{name}: no transaction header was detected.")

    headers = rows[header_index][1]
    date_column = _column(headers, ("date", "transaction date", "entry date"), name)
    description_column = _column(headers, DESCRIPTION_HEADERS, name)
    amount_column = _column(headers, AMOUNT_HEADERS, name)
    debit_column = _column(headers, ("debit",), name)
    credit_column = _column(headers, ("credit",), name)
    if date_column is None or description_column is None:
        raise ParseError("missing_column", f"{name}: date and description columns are required.")
    if amount_column is not None and (debit_column is not None or credit_column is not None):
        raise ParseError("ambiguous_columns", f"{name}: amount and debit/credit columns conflict.")
    split_amounts = debit_column is not None and credit_column is not None
    if not split_amounts and amount_column is None:
        raise ParseError("missing_column", f"{name}: an amount or debit/credit pair is required.")
    if not split_amounts and (debit_column is not None or credit_column is not None):
        raise ParseError("ambiguous_columns", f"{name}: both debit and credit columns are required.")

    descriptions = [str(_row_value(values, description_column)) for _number, values, _formulas in rows[header_index + 1:]
                    if not _empty(_row_value(values, description_column))]
    sheet_direction = None if split_amounts else _sheet_kind(name, descriptions)
    if not split_amounts and sheet_direction is None:
        raise ParseError("ambiguous_direction", f"{name}: income or expense direction is unclear.")

    for number, values, formulas in rows:
        raw = _raw_cells(values)
        if number <= rows[header_index][0]:
            excluded.append(ExcludedRow(name, number, raw, "header_or_prelude"))
            continue
        location = f"{name} row {number}"
        description_value = _row_value(values, description_column)
        description = str(description_value) if description_value is not None else ""
        if any(_normalize(value) in TOTAL_LABELS for value in values[:20] if isinstance(value, str)):
            excluded.append(ExcludedRow(name, number, raw, "total_or_balance"))
            continue
        critical_columns = {date_column, description_column, amount_column, debit_column, credit_column}
        if formulas.intersection(index for index in critical_columns if index is not None):
            raise ParseError("formula_value", f"{location}: a transaction field contains a formula; cached values may be stale.")

        date_value = _row_value(values, date_column)
        parsed_date = _date(date_value)
        if split_amounts:
            debit_value = _row_value(values, debit_column)
            credit_value = _row_value(values, credit_column)
            debit = _amount(debit_value, location)
            credit = _amount(credit_value, location)
            if debit and credit:
                raise ParseError("ambiguous_direction", f"{location}: both debit and credit have amounts.")
            signed_amount = debit if debit else credit if credit else None
            direction: Direction = "expense" if debit else "income"
        else:
            amount_value = _row_value(values, amount_column)
            signed_amount = _amount(amount_value, location)
            direction = sheet_direction or "expense"

        if signed_amount is None and parsed_date is None:
            financial_values = (debit_value, credit_value) if split_amounts else (amount_value,)
            if not _empty(date_value) or any(not _empty(value) for value in financial_values):
                raise ParseError("invalid_transaction", f"{location}: date or amount is invalid.")
            excluded.append(ExcludedRow(name, number, raw, "non_transaction"))
            continue
        if signed_amount is None:
            raise ParseError("missing_amount", f"{location}: dated entry has no usable amount.")
        if parsed_date is None:
            raise ParseError("invalid_date", f"{location}: amount has no valid date.")
        if not description.strip():
            raise ParseError("missing_description", f"{location}: amount has no description.")
        if len(description) > 2000:
            raise ParseError("description_length", f"{location}: description exceeds 2,000 characters.")
        if signed_amount == 0:
            excluded.append(ExcludedRow(name, number, raw, "zero_amount"))
            continue
        if signed_amount < 0:
            direction = "income" if direction == "expense" else "expense"
        amount = abs(signed_amount)
        included.append(ParsedRow(name, number, raw, parsed_date, description,
                                  direction, amount, _category(description, direction)))
    return included, excluded


def parse_workbook(content: bytes) -> ParsedWorkbook:
    """Parse one .xlsx deterministically or raise ParseError before any ledger write."""
    _check_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, ValueError, KeyError, OSError) as error:
        raise ParseError("invalid_workbook", "The workbook could not be read as .xlsx.") from error
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            raise ParseError("sheet_count", "The workbook has more than 50 sheets.")
        included: list[ParsedRow] = []
        excluded: list[ExcludedRow] = []
        populated = 0
        for sheet in workbook.worksheets:
            rows: list[tuple[int, tuple[object, ...], frozenset[int]]] = []
            for number, cells in enumerate(sheet.iter_rows(), start=1):
                if number > MAX_SCANNED_ROWS:
                    raise ParseError("row_span", f"{sheet.title}: row span exceeds the safety limit.")
                values = tuple(cell.value for cell in cells)
                if all(_empty(value) for value in values):
                    continue
                populated += 1
                if populated > MAX_POPULATED_ROWS:
                    raise ParseError("row_count", "The workbook has more than 100,000 populated rows.")
                formulas = frozenset(index for index, cell in enumerate(cells) if cell.data_type == "f")
                rows.append((number, values, formulas))
            if not rows:
                continue
            parsed, skipped = _parse_sheet(sheet.title, rows)
            included.extend(parsed)
            excluded.extend(skipped)
        if not included:
            raise ParseError("no_transactions", "No unambiguous transactions were found in the workbook.")
        return ParsedWorkbook(tuple(included), tuple(excluded), populated)
    except ParseError:
        raise
    except (ValueError, TypeError, OSError) as error:
        raise ParseError("invalid_workbook", "The workbook contains unreadable worksheet data.") from error
    finally:
        workbook.close()
